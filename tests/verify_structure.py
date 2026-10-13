"""Structural verification of generated DOCX/PPTX samples."""
from docx import Document
from pptx import Presentation

d = Document("tests/output/remote-work-policy-memo.docx")
paras = [p for p in d.paragraphs if p.text.strip()]
print("DOCX paragraphs:", len(paras), "| tables:", len(d.tables))
print("Title:", repr(paras[0].text))
print("Has TOC field:", any("TOC" in p._p.xml for p in d.paragraphs))
print("Header:", d.sections[0].header.paragraphs[0].text)
print("Footer text:", d.sections[0].footer.paragraphs[0].text)
for i, t in enumerate(d.tables):
    first_row = [c.text for c in t.rows[0].cells]
    print(f"  table {i}: rows={len(t.rows)} cols={len(t.columns)} head={first_row}")

print()
p = Presentation("tests/output/product-brief-deck.pptx")
print("PPTX slides:", len(p.slides))
for i, s in enumerate(p.slides, 1):
    title = s.shapes.title.text if s.shapes.title else "(no title)"
    has_footer = any(sh.has_text_frame and sh.text_frame.text.startswith("Atlas")
                     for sh in s.shapes)
    has_num = any(sh.has_text_frame and sh.text_frame.text.strip() == str(i)
                  for sh in s.shapes)
    print(f"  slide {i}: title={title!r} footer={has_footer} slide_num={has_num}")
